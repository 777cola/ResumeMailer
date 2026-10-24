"""ResumeMailer FastAPI 服务端"""
import base64
import io
import os
import re
import tempfile
import uuid
from datetime import date, datetime, timedelta
from pathlib import Path
from urllib.parse import quote

import openpyxl
import uvicorn
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Request
from fastapi.responses import HTMLResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles

from config import (EMAIL_PRESETS, load_config, save_config, EmailConfig, ATTACH_DIR,
                    LOG_FILE, sent_today, already_contacted, load_log,
                    DEFAULT_BODY_TEMPLATE, DEFAULT_SUBJECT_TEMPLATE)
from engine.excel_parser import generate_template, parse_workbook, parse_tsv
from engine.file_reader import file_kind, mime_of, read_text
from engine.ai_filler import (ai_chat, ai_fill_template, ai_extract_images,
                             ai_extract_files, ai_parse_resume, load_soul,
                             ai_generate_body, get_available_providers, ai_test)
from engine.sender import check_address, send_batch, send_test
from engine.bounce import scan_inbox
from engine import delivery
from engine.render import (render_email, unknown_placeholders, collect_vars,
                           build_context)
from engine import jobs

app = FastAPI(title="ResumeMailer · 简历投递助手")


@app.middleware("http")
async def _no_stale_static(request: Request, call_next):
    """升级后别让浏览器拿旧 JS/CSS——本地应用没有构建哈希，只能每次回源校验"""
    resp = await call_next(request)
    if request.url.path == "/" or request.url.path.startswith("/static/"):
        resp.headers["Cache-Control"] = "no-cache"
    return resp

STATIC_DIR = Path(__file__).parent / "static"


def _app_version() -> str:
    try:
        return (Path(__file__).parent / "VERSION").read_text(encoding="utf-8").strip()
    except OSError:
        return "dev"


APP_VERSION = _app_version()
app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")

# 公开部署时可设访问口令，避免别人读到你的配置 / 用你的实例发信
ACCESS_TOKEN = os.environ.get("RESUMEMAILER_TOKEN", "").strip()
MAX_ATTACH_MB = 25


@app.middleware("http")
async def guard(request: Request, call_next):
    if ACCESS_TOKEN and request.url.path.startswith("/api"):
        provided = request.headers.get("x-token") or request.query_params.get("token")
        if provided != ACCESS_TOKEN:
            return JSONResponse(status_code=401,
                                content={"ok": False, "need_token": True,
                                         "error": "需要访问口令（RESUMEMAILER_TOKEN）"})
    return await call_next(request)


@app.get("/", response_class=HTMLResponse)
def index():
    return (STATIC_DIR / "index.html").read_text(encoding="utf-8")


@app.get("/healthz")
def healthz():
    return {"ok": True}


# ── 上手引导 ──

DEMO_RECIPIENTS = [
    {"seq": 1, "company": "示例科技（虚构）", "email": "hr@example.com", "subject": "",
     "vars": {"岗位": "数据分析实习生", "HR": "张女士"}},
    {"seq": 2, "company": "虚构资本（虚构）", "email": "campus@example.com", "subject": "",
     "vars": {"岗位": "量化研究实习生", "HR": "李先生"}},
    {"seq": 3, "company": "样例制造（虚构）", "email": "zhaopin@example.com",
     "subject": "实习申请-{姓名}-{岗位}", "vars": {"岗位": "运营实习生", "HR": "王女士"}},
]

DEMO_BODY = """尊敬的{HR}：

您好！我是{学校}{专业}的{姓名}，看到贵司在招{岗位}，附上简历请查收。

电话/微信：{电话}
邮箱：{邮箱}"""


@app.get("/api/onboarding")
def onboarding():
    """上手清单的真实状态，前端据此打勾"""
    cfg = load_config()
    logs = load_log()
    sent = [e for e in logs if e.get("mode") == "send" and e.get("success")]
    attachments = [p for p in cfg.last_attachments if Path(p).is_file()]
    steps = [
        {"key": "email", "title": "配置发件邮箱",
         "done": bool(cfg.email.email and cfg.email.auth_code),
         "hint": "QQ 邮箱：设置 → 账户 → 生成授权码（不是登录密码）",
         "tab": "config", "detail": cfg.email.email or ""},
        {"key": "attachment", "title": "上传简历附件",
         "done": bool(attachments),
         "hint": "中文文件名不会乱码；配了 AI 还能从简历自动填个人信息",
         "tab": "attachments",
         "detail": f"{len(attachments)} 个文件" if attachments else ""},
        {"key": "body", "title": "确认正文与主题",
         "done": bool((cfg.body_template or "").strip() and (cfg.subject_template or "").strip()),
         "hint": "变量写成 {姓名} {公司名} {岗位}，发送时逐封替换",
         "tab": "body", "detail": ""},
        {"key": "recipients", "title": "导入收件人",
         "done": bool(cfg.last_recipients),
         "hint": "下载 Excel 模板填好传回来，或者把截图、PDF、Word 丢给 AI 提取",
         "tab": "recipients",
         "detail": f"{len(cfg.last_recipients)} 位" if cfg.last_recipients else ""},
        {"key": "test", "title": "先发给自己",
         "done": bool(cfg.last_test_at),
         "hint": "确认排版、附件名、主题都对，再群发",
         "tab": "send", "detail": cfg.last_test_at},
        {"key": "sent", "title": "开始投递",
         "done": bool(sent),
         "hint": "建议先「全部存草稿」，在网页邮箱里检查后再发",
         "tab": "send", "detail": f"已成功 {len(sent)} 封" if sent else ""},
    ]
    return {"steps": steps, "first_run": not cfg.email.email, "sent_total": len(sent)}


DEMO_USER = {"name": "张小明", "school": "示例大学", "major": "金融学",
             "grade": "27届", "phone": "13800000000", "email": "zhang@example.com"}


@app.post("/api/demo/load")
def demo_load():
    """载入三家公司示例，让人不填表也能完整体验一遍"""
    cfg = load_config()
    cfg.last_recipients = [dict(r, vars=dict(r["vars"])) for r in DEMO_RECIPIENTS]
    user_filled = []
    for key, value in DEMO_USER.items():
        if not getattr(cfg.user, key, ""):
            setattr(cfg.user, key, value)
            user_filled.append(key)
    body_set = False
    if not (cfg.body_template or "").strip() or cfg.body_template == DEFAULT_BODY_TEMPLATE:
        cfg.body_template = DEMO_BODY
        body_set = True
    if not (cfg.subject_template or "").strip() or cfg.subject_template == DEFAULT_SUBJECT_TEMPLATE:
        cfg.subject_template = "实习申请-{姓名}-{学校}-{岗位}"
        body_set = True
    save_config(cfg)
    return {"ok": True, "recipients": cfg.last_recipients, "body_set": body_set,
            "user_filled": user_filled, "variables": ["岗位", "HR"]}


# ── 配置 ──

def _config_payload() -> dict:
    cfg = load_config()
    return {
        "email": {
            "email": cfg.email.email,
            "smtp_host": cfg.email.smtp_host,
            "smtp_port": cfg.email.smtp_port,
            "imap_host": cfg.email.imap_host,
            "imap_port": cfg.email.imap_port,
            "use_ssl": cfg.email.use_ssl,
            "auth_code": bool(cfg.email.auth_code),   # 只返回是否存在
            "sender_name": cfg.email.sender_name,
        },
        "ai": {
            "provider": cfg.ai.provider,
            "api_key": bool(cfg.ai.api_key),
            "model": cfg.ai.model,
            "base_url": cfg.ai.base_url,
            "configured": bool(cfg.ai.provider and cfg.ai.api_key and cfg.ai.model),
        },
        "user": {
            "name": cfg.user.name, "school": cfg.user.school, "major": cfg.user.major,
            "grade": cfg.user.grade, "phone": cfg.user.phone, "email": cfg.user.email,
        },
        "body_template": cfg.body_template,
        "subject_template": cfg.subject_template,
        "send_rate": cfg.send_rate,
        "daily_limit": cfg.daily_limit,
        "cc_self": cfg.cc_self,
        "last_attachments": cfg.last_attachments,
        "last_recipients": cfg.last_recipients,
        "ai_providers": get_available_providers(),
        "email_presets": EMAIL_PRESETS,
        "today": {"sent": sent_today(), "date": date.today().isoformat()},
        "token_required": bool(ACCESS_TOKEN),
        "version": APP_VERSION,
        "running_job": jobs.active_job_id(),
        "limits": {"max_attach_mb": MAX_ATTACH_MB},
    }


@app.get("/api/config")
def get_config():
    return _config_payload()


@app.post("/api/config/email")
def save_email_config(data: dict):
    cfg = load_config()
    for key in ("email", "smtp_host", "smtp_port", "imap_host", "imap_port",
                "use_ssl", "sender_name"):
        if key in data and data[key] is not None:
            setattr(cfg.email, key, data[key])
    if data.get("auth_code"):
        cfg.email.auth_code = str(data["auth_code"]).strip()
    save_config(cfg)
    return {"ok": True}


@app.post("/api/config/ai")
def save_ai_config(data: dict):
    cfg = load_config()
    if data.get("provider"):
        cfg.ai.provider = data["provider"]
    if data.get("model"):
        cfg.ai.model = data["model"]
    if "base_url" in data:
        cfg.ai.base_url = str(data.get("base_url") or "").strip()
    if data.get("api_key"):
        cfg.ai.api_key = str(data["api_key"]).strip()
    save_config(cfg)
    return {"ok": True}


@app.post("/api/config/settings")
def save_settings(data: dict):
    """正文 / 主题 / 速度 / 每日上限 / 抄送自己"""
    cfg = load_config()
    if "body_template" in data:
        cfg.body_template = data["body_template"] or ""
    if "subject_template" in data:
        cfg.subject_template = data["subject_template"] or ""
    if "send_rate" in data:
        try:
            cfg.send_rate = max(1, min(120, int(data["send_rate"])))
        except (TypeError, ValueError):
            pass
    if "daily_limit" in data:
        try:
            cfg.daily_limit = max(0, min(1000, int(data["daily_limit"])))
        except (TypeError, ValueError):
            pass
    if "cc_self" in data:
        cfg.cc_self = bool(data["cc_self"])
    save_config(cfg)
    return {"ok": True}


@app.post("/api/config/user")
def save_user_info(data: dict):
    cfg = load_config()
    for key in ("name", "school", "major", "grade", "phone", "email"):
        if key in data:
            setattr(cfg.user, key, str(data[key] or "").strip())
    save_config(cfg)
    return {"ok": True}


# ── 连接测试 ──

@app.post("/api/email/test")
def test_email(data: dict):
    cfg = load_config()
    ecfg = EmailConfig(
        email=data.get("email") or cfg.email.email,
        smtp_host=data.get("smtp_host") or cfg.email.smtp_host,
        smtp_port=int(data.get("smtp_port") or cfg.email.smtp_port),
        imap_host=data.get("imap_host") or cfg.email.imap_host,
        imap_port=int(data.get("imap_port") or cfg.email.imap_port),
        use_ssl=bool(data.get("use_ssl", cfg.email.use_ssl)),
        auth_code=data.get("auth_code") or cfg.email.auth_code,
    )
    if not ecfg.email or not ecfg.auth_code:
        return JSONResponse(status_code=400,
                            content={"ok": False, "msg": "邮箱地址或授权码为空"})
    ok, msg = test_connection(ecfg)
    return {"ok": ok, "msg": msg} if ok else JSONResponse(
        status_code=400, content={"ok": False, "msg": msg})


# ── 模板 / 解析 ──

@app.get("/api/template/download")
def download_template():
    content = generate_template()
    fname = quote("ResumeMailer收件人模板.xlsx")
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{fname}"})


def _store_recipients(recipients: list):
    cfg = load_config()
    cfg.last_recipients = recipients
    save_config(cfg)


@app.post("/api/template/parse")
async def parse_excel(file: UploadFile = File(...)):
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(400, "文件超过 10MB，请确认是不是选错了文件")
    result = parse_workbook(content, file.filename or "")
    recipients = [r.to_dict() for r in result.recipients]
    if recipients:
        _store_recipients(recipients)
    return {"recipients": recipients, "errors": result.errors,
            "warnings": result.warnings, "total": len(recipients),
            "variables": collect_vars(recipients)}


@app.post("/api/recipients/parse-text")
def parse_text(data: dict):
    result = parse_tsv(data.get("text", ""))
    recipients = [r.to_dict() for r in result.recipients]
    if recipients:
        _store_recipients(recipients)
    return {"recipients": recipients, "errors": result.errors,
            "warnings": result.warnings, "total": len(recipients),
            "variables": collect_vars(recipients)}


@app.get("/api/recipients/last")
def last_recipients():
    cfg = load_config()
    return {"recipients": cfg.last_recipients, "variables": collect_vars(cfg.last_recipients)}


@app.post("/api/recipients/save")
def save_recipients(data: dict):
    _store_recipients(data.get("recipients") or [])
    return {"ok": True}


# ── 预览 / 预检 ──

@app.post("/api/preview")
def preview(data: dict):
    cfg = load_config()
    recipients = _requested_recipients(data, cfg)
    body = data.get("body", cfg.body_template)
    subject = data.get("subject", cfg.subject_template)
    user = data.get("user") or {
        "姓名": cfg.user.name, "学校": cfg.user.school, "专业": cfg.user.major,
        "年级": cfg.user.grade, "电话": cfg.user.phone, "邮箱": cfg.user.email,
    }
    limit = max(1, min(50, int(data.get("limit") or 5)))

    items = []
    for r in recipients[:limit]:
        rendered = render_email(r, user, body, subject)
        items.append({
            "seq": r.get("seq"), "company": r.get("company", ""),
            "email": r.get("email", ""), "subject": rendered["subject"],
            "body": rendered["body"], "html": rendered["html"],
            "vars": r.get("vars") or {},
            "error": check_address(r.get("email", "")),
        })

    invalid = [{"company": r.get("company", ""), "email": r.get("email", ""),
                "reason": check_address(r.get("email", ""))}
               for r in recipients if check_address(r.get("email", ""))]

    unknown = set()
    empty = set()
    for r in recipients:
        ctx = build_context(r, user)
        own_subject = (r.get("subject") or "") or (subject or "")
        unknown |= set(unknown_placeholders(own_subject, ctx))
        unknown |= set(unknown_placeholders(body or "", ctx))
        for name in set(re.findall(r"\{([^{}\r\n]{1,24})\}", own_subject + (body or ""))):
            if name in ctx and not str(ctx[name]).strip():
                empty.add(name)

    dead = delivery.dead_emails()
    dead_hits, seen_dead = [], set()
    for r in recipients:
        addr = str(r.get("email", "")).lower()
        if addr in dead and addr not in seen_dead:
            seen_dead.add(addr)
            dead_hits.append({"email": addr, "company": r.get("company", ""),
                              "reason": dead[addr]})

    contacted = already_contacted()
    dup_in_batch, contacted_hits, seen = [], [], set()
    for r in recipients:
        addr = str(r.get("email", "")).lower()
        if addr in seen:
            dup_in_batch.append(addr)
        seen.add(addr)
        if addr in contacted:
            contacted_hits.append({"email": addr, "company": r.get("company", ""),
                                   "last": contacted[addr]})

    return {
        "items": items,
        "total": len(recipients),
        "unknown_variables": sorted(unknown),
        "empty_variables": sorted(empty),
        "missing_company": [r.get("email") for r in recipients if not r.get("company")],
        "duplicates_in_batch": dup_in_batch,
        "already_contacted": contacted_hits[:50],
        "already_contacted_count": len(contacted_hits),
        "dead_addresses": dead_hits[:50],
        "dead_address_count": len(dead_hits),
        "invalid_addresses": invalid[:5],
        "invalid_address_count": len(invalid),
        "variables": collect_vars(recipients),
        "today": {"sent": sent_today(), "limit": cfg.daily_limit},
        "attachments": [Path(p).name for p in (cfg.last_attachments or [])],
    }


# ── AI ──

def _ai_params(data: dict):
    cfg = load_config()
    provider = data.get("provider") or cfg.ai.provider
    model = data.get("model") or cfg.ai.model
    key = data.get("api_key") or cfg.ai.api_key
    base_url = data.get("base_url") or cfg.ai.base_url
    if not provider or not key or not model:
        raise HTTPException(400, "还没配 AI：在「收件人 → AI 输入」里选一个提供商、粘贴 Key 就能用")
    if provider == "自定义（OpenAI 兼容）" and not base_url:
        raise HTTPException(400, "自定义提供商需要填 Base URL（形如 https://xxx/v1）")
    return provider, key, model, base_url


@app.post("/api/ai/test")
def ai_test_endpoint(data: dict):
    cfg = load_config()
    provider = data.get("provider") or cfg.ai.provider
    model = data.get("model") or cfg.ai.model
    key = data.get("api_key") or cfg.ai.api_key
    base_url = data.get("base_url") or cfg.ai.base_url
    if not provider or not model:
        raise HTTPException(400, "先选提供商和模型")
    if not key:
        raise HTTPException(400, "先粘贴 API Key")
    try:
        out = ai_test(provider, key, model, base_url=base_url)
        return {"ok": True, **out, "provider": provider, "model": model}
    except Exception as e:
        return JSONResponse(status_code=400, content={"ok": False, "error": str(e)})


@app.post("/api/ai/chat")
def ai_chat_endpoint(data: dict):
    provider, key, model, base_url = _ai_params(data)
    messages = data.get("messages") or []
    if not messages:
        raise HTTPException(400, "消息不能为空")
    try:
        return {"ok": True, "content": ai_chat(provider, key, model, messages,
                                               base_url=base_url)}
    except Exception as e:
        return JSONResponse(status_code=400, content={"ok": False, "error": str(e)})


@app.post("/api/ai/fill-template")
def ai_fill(data: dict):
    provider, key, model, base_url = _ai_params(data)
    text = (data.get("input") or "").strip()
    if not text:
        raise HTTPException(400, "输入内容不能为空")
    try:
        rows = ai_fill_template(provider, key, model, text, base_url=base_url)
        recipients = []
        for i, r in enumerate(rows if isinstance(rows, list) else []):
            recipients.append({
                "seq": i + 1,
                "company": r.get("company", ""),
                "email": r.get("email", ""),
                "subject": r.get("subject", ""),
                "vars": r.get("vars") or {},
            })
        if recipients:
            _store_recipients(recipients)
        return {"ok": True, "recipients": recipients,
                "variables": collect_vars(recipients)}
    except Exception as e:
        return JSONResponse(status_code=400, content={"ok": False, "error": str(e)})


@app.post("/api/ai/parse-files")
async def ai_parse_files(files: list[UploadFile] = File(...),
                         provider: str = Form(""), model: str = Form(""),
                         api_key: str = Form(""), hint: str = Form("")):
    """任何格式的材料（截图/PDF/Word/Excel/文本）→ 收件人名单"""
    cfg = load_config()
    provider = provider or cfg.ai.provider
    model = model or cfg.ai.model
    key = api_key or cfg.ai.api_key
    base_url = cfg.ai.base_url
    if not provider or not model or not key:
        raise HTTPException(400, "先在这个页面的「1. 接上 AI」里选一个提供商、粘贴 API Key")

    blocks, used, skipped = [], [], []
    for f in files[:6]:
        name = Path(f.filename or "材料").name
        data = await f.read()
        if len(data) > 12 * 1024 * 1024:
            skipped.append(f"{name} 超过 12MB")
            continue
        kind = file_kind(Path(name))
        if kind == "image":
            mime = (f.content_type or "").lower()
            if not mime.startswith("image/"):
                mime = mime_of(Path(name))
            blocks.append({"kind": "image", "name": name,
                           "data_url": f"data:{mime};base64,{base64.b64encode(data).decode()}"})
            used.append(f"{name}（图片）")
            continue
        # 非图片：落到临时文件再按后缀读，避免把路径当输入
        tmp = Path(tempfile.gettempdir()) / f"rm_{uuid.uuid4().hex}{Path(name).suffix.lower()}"
        tmp.write_bytes(data)
        try:
            text = read_text(tmp)
        except ValueError as e:
            skipped.append(str(e))
            continue
        finally:
            try:
                tmp.unlink()
            except OSError:
                pass
        if not text.strip():
            skipped.append(f"{name} 里没有可读文字（扫描件请先导出成 PDF/Word，或截图给我看）")
            continue
        blocks.append({"kind": "text", "name": name, "text": text})
        used.append(f"{name}（{len(text)} 字）")

    if not blocks:
        raise HTTPException(400, "这些文件里没有可以读的内容：" + "；".join(skipped or ["未知格式"]))

    try:
        rows = ai_extract_files(provider, key, model, blocks, hint=hint.strip(), base_url=base_url)
    except ValueError as e:
        return JSONResponse(status_code=400, content={"ok": False, "error": str(e)})
    except Exception as e:
        return JSONResponse(status_code=400, content={"ok": False, "error": str(e)})

    recipients = [{"seq": i + 1, "company": r.get("company", ""),
                   "email": r.get("email", ""), "subject": r.get("subject", ""),
                   "vars": r.get("vars") or {}, "note": r.get("note", "")}
                  for i, r in enumerate(rows)]
    return {"ok": True, "recipients": recipients, "total": len(recipients),
            "files_used": used, "files_skipped": skipped,
            "model": f"{provider} · {model}"}


@app.post("/api/ai/resume-fill")
def ai_resume_fill(data: dict):
    """读一份简历附件，抽出个人信息填正文"""
    cfg = load_config()
    provider = data.get("provider") or cfg.ai.provider
    model = data.get("model") or cfg.ai.model
    key = data.get("api_key") or cfg.ai.api_key
    base_url = data.get("base_url") or cfg.ai.base_url
    if not provider or not model or not key:
        raise HTTPException(400, "还没配 AI：在「收件人 → AI 输入」里选提供商、粘贴 Key 后就能自动读简历")

    path = data.get("filename") or ""
    if not path:
        return JSONResponse(status_code=400, content={"ok": False, "error": "先选一份简历附件"})
    try:
        safe = _safe_attachment(path)
        text = read_text(Path(safe))
    except HTTPException:
        raise
    except ValueError as e:
        return JSONResponse(status_code=400, content={"ok": False, "error": str(e)})

    if not text.strip():
        return JSONResponse(status_code=400,
                            content={"ok": False, "error": "这份文件读不出文字（扫描版简历请导出成 PDF/Word）"})
    try:
        fields = ai_parse_resume(provider, key, model, text, base_url=base_url)
    except Exception as e:
        return JSONResponse(status_code=400, content={"ok": False, "error": str(e)})
    return {"ok": True, "fields": fields, "filename": Path(safe).name}


@app.get("/api/ai/soul")
def ai_soul(name: str = "recipient"):
    """把 AI 的角色设定原文给界面看，用户知道它按什么规则在干活"""
    if name not in ("recipient", "resume", "body"):
        raise HTTPException(400, "未知的 soul")
    text = load_soul(name)
    if not text:
        raise HTTPException(404, "没有找到 soul 文件")
    return {"ok": True, "name": name, "content": text}


@app.post("/api/ai/parse-images")
async def ai_parse_images(files: list[UploadFile] = File(...),
                          provider: str = Form(""), model: str = Form(""),
                          api_key: str = Form(""), hint: str = Form("")):
    """上传截图/照片，AI 提取公司、邮箱、岗位"""
    cfg = load_config()
    provider = provider or cfg.ai.provider
    model = model or cfg.ai.model
    key = api_key or cfg.ai.api_key
    base_url = cfg.ai.base_url
    if not provider or not model or not key:
        raise HTTPException(400, "先选一个支持看图的模型并填 API Key（GPT-4o / Claude / qwen-vl-max / glm-4v-plus）")

    images = []
    for f in files[:6]:
        data = await f.read()
        if len(data) > 8 * 1024 * 1024:
            raise HTTPException(400, f"{f.filename} 超过 8MB，压缩后再传")
        mime = (f.content_type or "").lower()
        if not mime.startswith("image/"):
            raise HTTPException(400, f"{f.filename} 不是图片")
        images.append((data, mime))
    if not images:
        raise HTTPException(400, "没有收到图片")

    try:
        rows = ai_extract_images(provider, key, model, images, hint=hint.strip(),
                                 base_url=base_url)
        recipients = [{"seq": i + 1, "company": r.get("company", ""),
                       "email": r.get("email", ""), "subject": r.get("subject", ""),
                       "vars": r.get("vars") or {}} for i, r in enumerate(rows)]
        if recipients:
            _store_recipients(recipients)
        return {"ok": True, "recipients": recipients,
                "variables": collect_vars(recipients)}
    except Exception as e:
        return JSONResponse(status_code=400, content={"ok": False, "error": str(e)})


@app.post("/api/ai/generate-body")
def ai_gen_body(data: dict):
    provider, key, model, base_url = _ai_params(data)
    instruction = (data.get("instruction") or "").strip()
    if not instruction:
        raise HTTPException(400, "请输入修改要求")
    cfg = load_config()
    try:
        content = ai_generate_body(provider, key, model, instruction,
                                   current_body=data.get("current_body", cfg.body_template),
                                   user={"姓名": cfg.user.name, "学校": cfg.user.school,
                                         "专业": cfg.user.major, "年级": cfg.user.grade},
                                   base_url=base_url)
        return {"ok": True, "content": content}
    except Exception as e:
        return JSONResponse(status_code=400, content={"ok": False, "error": str(e)})


# ── 附件 ──

def _safe_attachment(path: str) -> str:
    """只允许使用上传目录里的文件，防止公开部署时被读任意路径"""
    p = Path(path).resolve()
    root = ATTACH_DIR.resolve()
    if p.parent != root:
        raise HTTPException(400, f"附件路径非法：{Path(path).name}")
    if not p.is_file():
        raise HTTPException(400, f"附件不存在：{Path(path).name}")
    return str(p)


@app.get("/api/attachments")
def list_attachments():
    files = []
    if ATTACH_DIR.exists():
        for f in sorted(ATTACH_DIR.iterdir(), key=lambda x: x.stat().st_mtime, reverse=True):
            if f.is_file():
                files.append({"name": f.name, "path": str(f), "size": f.stat().st_size})
    return {"files": files}


@app.post("/api/attachments/upload")
async def upload_attachment(file: UploadFile = File(...)):
    ATTACH_DIR.mkdir(parents=True, exist_ok=True)
    name = Path(file.filename or "attachment").name
    save_path = ATTACH_DIR / name
    if save_path.exists():
        save_path = ATTACH_DIR / f"{save_path.stem}_{datetime.now():%H%M%S}{save_path.suffix}"
    content = await file.read()
    if len(content) > MAX_ATTACH_MB * 1024 * 1024:
        raise HTTPException(400, f"单个附件不能超过 {MAX_ATTACH_MB}MB")
    save_path.write_bytes(content)

    cfg = load_config()
    if str(save_path) not in cfg.last_attachments:
        cfg.last_attachments.append(str(save_path))
        save_config(cfg)
    return {"ok": True, "file": {"name": save_path.name, "path": str(save_path),
                                 "size": len(content)}}


@app.delete("/api/attachments/{filename}")
def delete_attachment(filename: str):
    filepath = ATTACH_DIR / Path(filename).name
    if filepath.is_file():
        filepath.unlink()
        cfg = load_config()
        cfg.last_attachments = [p for p in cfg.last_attachments if Path(p).name != filepath.name]
        save_config(cfg)
    return {"ok": True}


# ── 发送 ──

def _requested_recipients(data: dict, cfg) -> list:
    """显式传空列表 = 没有收件人；没传字段才沿用上次的"""
    raw = data["recipients"] if "recipients" in data else cfg.last_recipients
    return raw or []


def _requested_attachments(data: dict, cfg) -> list:
    """显式传空列表 = 不发附件；没传字段 = 沿用上次选择"""
    raw = data["attachments"] if "attachments" in data else cfg.last_attachments
    return [_safe_attachment(p) for p in (raw or [])]


def _check_ready(cfg) -> None:
    if not cfg.email.email or not cfg.email.auth_code:
        raise HTTPException(400, "请先在「邮箱配置」页填写邮箱和授权码")


@app.post("/api/send/test")
def send_to_self(data: dict):
    """先发给自己一封，确认排版和附件名"""
    cfg = load_config()
    _check_ready(cfg)
    recipients = _requested_recipients(data, cfg)
    sample = recipients[0] if recipients else {"company": "示例公司", "email": "", "vars": {}}
    user = data.get("user") or {"姓名": cfg.user.name, "学校": cfg.user.school,
                                "专业": cfg.user.major, "年级": cfg.user.grade,
                                "电话": cfg.user.phone, "邮箱": cfg.user.email}
    rendered = render_email(sample, user, data.get("body", cfg.body_template),
                            data.get("subject", cfg.subject_template))
    attachments = _requested_attachments(data, cfg)
    try:
        addr = send_test(cfg.email, rendered["subject"], rendered["body"],
                         rendered["html"], attachments)
    except Exception as e:
        return JSONResponse(status_code=400, content={"ok": False, "error": str(e)})
    cfg.last_test_at = datetime.now().isoformat(timespec="seconds")
    save_config(cfg)
    return {"ok": True, "msg": f"已发送到 {addr}，请到收件箱确认排版与附件名"}


RECENT_DAYS = 30


def recent_logs(days=RECENT_DAYS, logs=None):
    """只取最近 N 天的记录。时间解析不了的条目保留，宁可多显示不可丢。"""
    src = load_log() if logs is None else logs
    cutoff = datetime.now() - timedelta(days=days)
    out = []
    for e in src:
        raw = str(e.get("time", "") or "").strip()
        try:
            when = datetime.fromisoformat(raw.replace("Z", ""))
        except ValueError:
            out.append(e)
            continue
        if when >= cutoff:
            out.append(e)
    return out


def slice_batch(recipients, skip_sent=False, limit=0, logs=None):
    """先发第一封 / 只发一部分：返回 (本次要发的列表, 因为已发过而跳过的数量)

    skip_sent 只认日志里 success 且 status=sent 的记录；存过草稿的不算已发（草稿另发会重复，
    但那是用户自己的选择，这里不替他判断）。
    """
    out = list(recipients or [])
    sent_before = 0
    if skip_sent:
        src = load_log() if logs is None else logs
        done = {str(e.get("email", "")).lower() for e in src
                if e.get("success") and e.get("status") == "sent"}
        keep = [r for r in out if str(r.get("email", "")).lower() not in done]
        sent_before = len(out) - len(keep)
        out = keep
    if limit and limit > 0:
        out = out[:limit]
    return out, sent_before


@app.post("/api/send/start")
def send_start(data: dict):
    cfg = load_config()

    recipients = _requested_recipients(data, cfg)
    if not recipients:
        raise HTTPException(400, "收件人列表为空")
    if len(recipients) > 2000:
        raise HTTPException(400, "单批最多 2000 封，请拆分发送")
    full_recipients = list(recipients)   # 用户完整名单；limit / skip_sent 只影响本次任务

    # 先发第一封 / 只发一部分：limit 截断，skip_sent 跳过已成功投递过的地址
    limit = max(0, int(data.get("limit") or 0))
    recipients, sent_before = slice_batch(recipients, bool(data.get("skip_sent")), limit)
    if not recipients:
        raise HTTPException(400, "这批收件人已经全部发过了")

    mode = data.get("send_mode", "draft")
    if mode not in ("draft", "send", "dry"):
        raise HTTPException(400, "send_mode 只能是 draft / send / dry")
    if mode != "dry":            # 试跑不连服务器，不需要邮箱
        _check_ready(cfg)

    attachments = _requested_attachments(data, cfg)
    body = data.get("body", cfg.body_template)
    subject = data.get("subject", cfg.subject_template)
    rate = int(data.get("rate") or cfg.send_rate or 20)
    daily_limit = int(data.get("daily_limit", cfg.daily_limit) or 0)
    cc_self = bool(data.get("cc_self", cfg.cc_self))

    # 之前退信过的地址默认跳过（可在请求里传 include_dead=true 强制重试）
    include_dead = bool(data.get("include_dead"))
    dead = {} if include_dead else delivery.dead_emails()
    job_recipients, dead_skipped = [], 0
    for r in recipients:
        item = dict(r)
        addr = str(item.get("email", "")).lower()
        if addr in dead:
            reason = dead[addr][:60] if dead[addr] else "此前投递被退回"
            item["skip_reason"] = f"退信地址，已自动跳过（{reason}）"
            dead_skipped += 1
        job_recipients.append(item)

    cfg.last_recipients = full_recipients  # 原样保存，不把 skip_reason 写进配置
    cfg.last_attachments = attachments
    cfg.body_template = body
    cfg.subject_template = subject
    save_config(cfg)

    user = {"姓名": cfg.user.name, "学校": cfg.user.school, "专业": cfg.user.major,
            "年级": cfg.user.grade, "电话": cfg.user.phone, "邮箱": cfg.user.email}
    try:
        job_id = jobs.start_job(
            cfg_email=cfg.email, recipients=job_recipients, body=body,
            subject_template=subject, attachments=attachments, mode=mode,
            rate=rate, user_info=user, daily_limit=daily_limit, cc_self=cc_self)
    except RuntimeError as e:
        raise HTTPException(409, str(e))
    return {"ok": True, "job_id": job_id, "total": len(recipients),
            "dead_skipped": dead_skipped, "already_sent": sent_before, "limit": limit,
            "skipped_by_limit": max(0, len(recipients) - max(0, daily_limit - sent_today())) if daily_limit else 0}


@app.get("/api/send/status/{job_id}")
def send_status(job_id: str):
    job = jobs.get_job(job_id)
    if not job:
        raise HTTPException(404, "任务不存在（服务重启后任务不保留）")
    return job


@app.post("/api/send/cancel/{job_id}")
def send_cancel(job_id: str):
    ok = jobs.cancel_job(job_id)
    return {"ok": ok, "msg": "已请求取消，当前这封发完即停止" if ok else "任务已结束"}


# ── 日志 / 导出 ──

@app.get("/api/logs")
def get_logs():
    logs = recent_logs()
    return {"logs": logs[-300:], "days": RECENT_DAYS, "total": len(logs)}


@app.delete("/api/logs")
def clear_logs():
    if LOG_FILE.exists():
        LOG_FILE.unlink()
    return {"ok": True}


@app.get("/api/contacts")
def contacts():
    rows = recent_logs()
    view = delivery.contacts_view(rows)
    counts = {"sent": 0, "replied": 0, "bounced": 0}
    for row in view:
        counts[row["status"]] = counts.get(row["status"], 0) + 1
    return {"contacts": view[:500], "counts": counts, "last_scan": delivery.last_scan(),
            "days": RECENT_DAYS, "log_total": len(rows)}


@app.post("/api/inbox/scan")
def inbox_scan(data: dict = None):
    """扫收件箱：Mailer-Daemon 退信 + 对方回信。只读，不改已读状态。"""
    cfg = load_config()
    if not (cfg.email.email and cfg.email.auth_code and cfg.email.imap_host):
        raise HTTPException(400, "需要先配置邮箱，并开启 IMAP（QQ/163 都要在设置里手动开）")
    days = max(1, min(90, int((data or {}).get("days", 14) or 14)))
    contacted = set(already_contacted().keys())
    msgids = {}
    for e in load_log():
        mid = str(e.get("message_id", "") or "").strip()
        if mid and e.get("success"):
            msgids[mid] = str(e.get("email", ""))
            msgids[mid.strip("<>")] = str(e.get("email", ""))
    try:
        out = scan_inbox(cfg.email, contacted, days=days, max_msgs=300, msgids=msgids)
    except Exception as e:
        return JSONResponse(status_code=400, content={"ok": False, "error": f"扫描失败：{e}"})
    if not out["scanned"] and out["errors"]:
        return JSONResponse(status_code=400, content={"ok": False, "error": out["errors"][0]})
    stats = delivery.apply_scan(out["bounces"], out["replies"])
    return {"ok": True, "scanned": out["scanned"], "bounces": out["bounces"],
            "replies": out["replies"], "errors": out["errors"][:5], **stats}


@app.get("/api/report/export")
def export_report():
    logs = recent_logs()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"投递记录（最近{RECENT_DAYS}天）"
    headers = ["时间", "公司", "收件人邮箱", "邮件主题", "方式", "结果", "备注"]
    ws.append(headers)
    widths = [20, 28, 34, 44, 10, 8, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="1A4B8C")
    for e in logs:
        ws.append([
            e.get("time", ""), e.get("company", ""), e.get("email", ""),
            e.get("subject", ""),
            "直接发送" if e.get("mode") == "send" else "存草稿",
            "成功" if e.get("success") else "失败",
            e.get("error", ""),
        ])
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    fname = quote(f"ResumeMailer投递记录-最近{RECENT_DAYS}天-{date.today():%Y%m%d}.xlsx")
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{fname}"})


# ── 启动 ──

def start():
    import webbrowser
    port = int(os.environ.get("PORT", 8765))
    host = os.environ.get("HOST", "127.0.0.1")
    is_cloud = host == "0.0.0.0"
    print(f" ResumeMailer 已启动 → http://localhost:{port}")
    if is_cloud and not ACCESS_TOKEN:
        print("公网监听但未设置 RESUMEMAILER_TOKEN，任何人可访问你的配置与附件")
    if not is_cloud:
        webbrowser.open(f"http://localhost:{port}")
    uvicorn.run(app, host=host, port=port, log_level="info")
