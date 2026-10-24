"""离线自检 — 不联网、不发信，覆盖解析 / 渲染 / MIME / 限速 / 取消

用法：python tests/test_selftest.py
任何一条挂了，说明改坏了真实发送路径。
"""
import io
import smtplib
import sys
import threading
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl

import engine.sender as sender
from engine.excel_parser import parse_table, parse_tsv, generate_template, parse_workbook
from engine.render import render_email, unknown_placeholders, to_html


def _xlsx(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_basic_and_vars():
    data = _xlsx([
        ["公司名称", "收件人邮箱", "邮件主题", "岗位"],
        ["华泰资产", "HR@HTAsset.com", "实习申请-{姓名}-{岗位}", "量化研究"],
        ["中信证券", "hr@citics.com", "", "行研"],
    ])
    res = parse_workbook(data, "t.xlsx")
    assert len(res.recipients) == 2, res.errors
    r0 = res.recipients[0]
    assert r0.email == "hr@htasset.com", "邮箱应小写"
    assert r0.vars["岗位"] == "量化研究"
    assert r0.subject.startswith("实习申请")


def test_parse_multi_email_splits():
    data = _xlsx([
        ["公司", "邮箱"],
        ["某公司", "a@x.com；b@y.com、c@z.com"],
    ])
    res = parse_workbook(data, "t.xlsx")
    assert len(res.recipients) == 3, [r.email for r in res.recipients]
    assert len({r.email for r in res.recipients}) == 3


def test_parse_dedupes():
    data = _xlsx([
        ["公司", "邮箱"],
        ["A", "hr@same.com"],
        ["B", "hr@same.com"],
    ])
    res = parse_workbook(data, "t.xlsx")
    assert len(res.recipients) == 1
    assert any("重复" in e for e in res.errors)


def test_parse_header_aliases_and_offset_header():
    data = _xlsx([
        ["2026 秋招投递表"],
        ["单位", "投递邮箱", "标题"],
        ["某基金", "hr@fund.com", "应聘-研究"],
    ])
    res = parse_workbook(data, "t.xlsx")
    assert len(res.recipients) == 1
    assert res.recipients[0].company == "某基金"


def test_parse_tsv_paste():
    res = parse_tsv("公司\t邮箱\nA公司\ta@a.com\nB公司\tb@b.com")
    assert len(res.recipients) == 2


def test_template_downloads():
    content = generate_template()
    assert content[:2] == b"PK"
    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert "填写说明" in wb.sheetnames


def test_render_subject_and_body():
    user = {"姓名": "AutumnPants", "学校": "华东理工", "专业": "金融"}
    r = {"company": "某公司", "email": "a@b.com", "subject": "实习-{姓名}-{岗位}",
         "vars": {"岗位": "量化"}}
    out = render_email(r, user, "你好 {公司名}，我是{姓名}。", "实习申请-{姓名}")
    assert out["subject"] == "实习-AutumnPants-量化", out["subject"]
    assert out["body"] == "你好 某公司，我是AutumnPants。"
    assert "AutumnPants" in out["html"]


def test_render_unknown_kept():
    out = render_email({"company": "A", "email": "a@b.com"}, {},
                       "你好 {不存在的变量}", "{姓名}")
    assert "{不存在的变量}" in out["body"]
    assert unknown_placeholders(out["body"], {}) == ["不存在的变量"]


def test_html_escaping():
    html = to_html("a<b> & c\nd")
    assert "&lt;b&gt;" in html and "&amp;" in html and "<br>" in html


def test_message_headers():
    from config import EmailConfig
    cfg = EmailConfig(email="me@qq.com", sender_name="AutumnPants")
    msg = sender.build_message(cfg, "hr@corp.com", "实习申请-AutumnPants",
                               "正文第一行\n第二行", to_html("正文第一行\n第二行"),
                               [], cc_self=True)
    assert msg["Date"], "缺 Date 头"
    assert msg["Message-ID"], "缺 Message-ID 头"
    assert msg.get_content_type() == "multipart/alternative", msg.get_content_type()
    types = [p.get_content_type() for p in msg.walk()]
    assert "text/plain" in types and "text/html" in types
    # 中文主题 / 发件人名必须能被标准库解回来
    import email
    from email import policy
    parsed = email.message_from_bytes(msg.as_bytes(), policy=policy.default)
    assert parsed["Subject"] == "实习申请-AutumnPants", parsed["Subject"]
    assert "AutumnPants" in str(parsed["From"])
    assert parsed["Cc"] == "me@qq.com"


def test_attachment_name_roundtrip():
    import tempfile
    from config import EmailConfig
    tmp = Path(tempfile.mkdtemp()) / "AutumnPants-简历.pdf"
    tmp.write_bytes(b"%PDF-1.4 test")
    msg = sender.build_message(EmailConfig(email="me@qq.com"), "hr@corp.com",
                               "s", "b", "", [str(tmp)])
    import email
    from email import policy
    parsed = email.message_from_bytes(msg.as_bytes(), policy=policy.default)
    names = [p.get_filename() for p in parsed.walk() if p.get_filename()]
    assert names == ["AutumnPants-简历.pdf"], names


def test_imap_utf7():
    assert sender._imap_utf7("Drafts") == "Drafts"
    enc = sender._imap_utf7("草稿箱")
    assert enc.startswith("&") and enc.endswith("-"), enc


def test_send_batch_daily_limit_and_skip():
    class FakeLink:
        sent = []

        def __init__(self, cfg):
            pass

        def send(self, msg):
            FakeLink.sent.append(msg["To"])

        def close(self):
            pass

    orig = sender._SmtpLink
    sender._SmtpLink = FakeLink
    try:
        from config import EmailConfig
        recipients = [{"company": f"C{i}", "email": f"a{i}@x.com"} for i in range(5)]
        out = sender.send_batch(EmailConfig(email="me@qq.com"), recipients,
                                "hi {公司名}", "主题-{公司名}", [],
                                send_mode="send", rate=600, daily_limit=3)
        statuses = [r["status"] for r in out["results"]]
        assert statuses.count("sent") == 3, statuses
        assert statuses.count("skipped") == 2, statuses
        assert "上限" in out["stopped"]
        assert len(FakeLink.sent) == 3
    finally:
        sender._SmtpLink = orig


def test_send_batch_cancel():
    class FakeLink:
        def __init__(self, cfg):
            pass

        def send(self, msg):
            pass

        def close(self):
            pass

    orig = sender._SmtpLink
    sender._SmtpLink = FakeLink
    try:
        from config import EmailConfig
        cancel = threading.Event()
        recipients = [{"company": f"C{i}", "email": f"a{i}@x.com"} for i in range(4)]

        def on_result(r):
            if r["status"] == "sent":
                cancel.set()

        out = sender.send_batch(EmailConfig(email="me@qq.com"), recipients,
                                "hi", "s", [], send_mode="send", rate=600,
                                progress=on_result, cancel=cancel)
        assert out["results"][0]["status"] == "sent"
        assert out["results"][-1]["status"] == "skipped"
        assert out["stopped"] == "已取消"
    finally:
        sender._SmtpLink = orig


def test_dry_run_never_touches_network():
    """试跑只组装消息，不能创建任何连接，也不吃每日额度"""
    class Boom:
        def __init__(self, *a, **k):
            raise AssertionError("试跑不应该创建 SMTP/IMAP 连接")

    orig_link, orig_sink = sender._SmtpLink, sender._DraftSink
    sender._SmtpLink = Boom
    sender._DraftSink = Boom
    try:
        from config import EmailConfig
        recs = [{"company": f"C{i}", "email": f"a{i}@x.com"} for i in range(3)]
        out = sender.send_batch(EmailConfig(email="me@qq.com"), recs, "hi {公司名}",
                                "t-{公司名}", [], send_mode="dry", rate=60, daily_limit=1)
        assert [r["status"] for r in out["results"]] == ["dry", "dry", "dry"], out["results"]
        assert all(r["success"] for r in out["results"])
        assert out["stopped"] == ""
    finally:
        sender._SmtpLink, sender._DraftSink = orig_link, orig_sink


def test_job_counters_treat_dry_as_success(monkeypatch=None):
    import engine.jobs as jobs

    class FakeJob(dict):
        pass

    job = FakeJob(id="t", status="running", mode="dry", total=2, done=0, success=0,
                  failed=0, skipped=0, started_at="", finished_at="", error="",
                  stopped="", draft_box="", results=[], cancel=threading.Event())
    orig = jobs.send_batch
    rows = [
        {"company": "A", "email": "a@x.com", "status": "dry", "success": True, "error": "", "subject": "s"},
        {"company": "B", "email": "b@x.com", "status": "dry", "success": True, "error": "", "subject": "s"},
    ]

    def fake_send_batch(**kw):
        progress = kw.get("progress")
        for r in rows:
            if progress:
                progress(r)
        return {"results": rows, "stopped": "", "draft_box": ""}

    jobs.send_batch = fake_send_batch
    try:
        jobs._run(job, None, [], "b", "s", [], 60, {}, 0, False)
    finally:
        jobs.send_batch = orig
    assert job["success"] == 2 and job["failed"] == 0, dict(job)


def test_auth_error_stops_batch():
    class FakeLink:
        def __init__(self, cfg):
            pass

        def send(self, msg):
            raise smtplib.SMTPAuthenticationError(535, b"bad auth")

        def close(self):
            pass

    orig = sender._SmtpLink
    sender._SmtpLink = FakeLink
    try:
        from config import EmailConfig
        recipients = [{"company": "A", "email": "a@x.com"},
                      {"company": "B", "email": "b@x.com"}]
        out = sender.send_batch(EmailConfig(email="me@qq.com"), recipients,
                                "hi", "s", [], send_mode="send", rate=600)
        assert all(r["status"] != "sent" for r in out["results"])
        assert "认证" in out["stopped"], out["stopped"]
    finally:
        sender._SmtpLink = orig


def test_header_company_mail_is_email_column():
    """「公司邮箱」是邮箱列，不能因为含「公司」被认成公司名"""
    data = _xlsx([
        ["公司", "公司邮箱", "主题"],
        ["某基金", "hr@fund.com", "应聘"],
    ])
    res = parse_workbook(data, "t.xlsx")
    assert len(res.recipients) == 1, res.errors
    assert res.recipients[0].company == "某基金", res.recipients[0].company
    assert res.recipients[0].email == "hr@fund.com"


def test_split_display_name():
    data = _xlsx([["公司", "邮箱"], ["A", "王女士 <HR@A.com>, li@a.com"]])
    res = parse_workbook(data, "t.xlsx")
    assert [r.email for r in res.recipients] == ["hr@a.com", "li@a.com"]


def test_csv_gbk():
    raw = "公司,邮箱\n某券商,hr@broker.com\n".encode("gbk")
    res = parse_workbook(raw, "t.csv")
    assert len(res.recipients) == 1, res.errors
    assert res.recipients[0].company == "某券商"


def test_old_xls_rejected():
    res = parse_workbook(b"whatever", "old.xls")
    assert res.recipients == [] and "xls" in res.errors[0]


def test_draft_box_detection():
    """Gmail / 163 的草稿箱名字不同，必须按 LIST 结果挑"""
    class FakeConn:
        def list(self):
            return ("OK", [b'(\\HasNoChildren) "/" "Drafts"',
                           b'(\\HasNoChildren) "/" "&XfJT0ZAB-"'])

    sink = sender._DraftSink(type("C", (), {"imap_host": "x", "imap_port": 993,
                                            "email": "a@b.c", "auth_code": "x"})())
    sink.conn = FakeConn()
    boxes = sink._boxes()
    assert boxes[0] == "Drafts", boxes
    assert any(b.startswith("&") for b in boxes), boxes


def test_html_body_plain_fallback():
    out = render_email({"company": "A", "email": "a@b.com"}, {},
                       "<p>你好</p><p>第二段</p>", "主题")
    assert "你好" in out["body"] and "第二段" in out["body"]
    assert "<p>" not in out["body"], out["body"]
    assert out["html"].startswith("<p>你好</p>"), out["html"]


def test_draft_append_falls_back_to_second_folder():
    """第一个候选箱名不存在时，要接着试下一个，并记住成功的那个"""
    class FakeConn:
        def __init__(self):
            self.calls = []

        def list(self):
            return ("OK", [b'(\\HasNoChildren) "/" "Drafts"',
                           ('(\\HasNoChildren) "/" "%s"' % sender._imap_utf7("草稿箱")).encode()])

        def append(self, name, flags, date, raw):
            self.calls.append(name)
            return ("OK", [b"ok"]) if name.strip('"') == sender._imap_utf7("草稿箱") else ("NO", [b"nope"])

        def logout(self):
            pass

    from config import EmailConfig
    sink = sender._DraftSink(EmailConfig(email="a@b.c", auth_code="x"))
    sink.conn = FakeConn()
    msg = sender.build_message(EmailConfig(email="a@b.c"), "hr@c.com", "s", "b", "", [])
    sink.append(msg)
    assert sink.box == sender._imap_utf7("草稿箱"), sink.box
    assert len(sink.conn.calls) == 2, sink.conn.calls
    # 第二次应直接用缓存成功的箱名
    sink.append(msg)
    assert len(sink.conn.calls) == 3, sink.conn.calls


def test_empty_attachment_list_does_not_fall_back():
    """显式传空附件 = 不发附件，不能回退到上次的简历"""
    import server
    cfg = type("C", (), {"last_attachments": ["/tmp/should_not_be_used.pdf"]})()
    assert server._requested_attachments({"attachments": []}, cfg) == []
    # 缺失字段时才沿用上次选择，并且要校验文件确实存在
    try:
        server._requested_attachments({}, cfg)
    except Exception as e:
        assert "不存在" in str(getattr(e, "detail", e)) or "非法" in str(getattr(e, "detail", e))
    else:
        raise AssertionError("沿用的附件不存在时必须报错，而不是静默发一封没有附件的信")


DSN_RAW = """From: Mail Delivery Subsystem <MAILER-DAEMON@qq.com>
To: me@qq.com
Subject: Undelivered Mail Returned to Sender
Date: Wed, 21 Oct 2026 10:00:00 +0800
MIME-Version: 1.0
Content-Type: multipart/report; report-type=delivery-status; boundary="BOUND"

--BOUND
Content-Type: text/plain; charset=utf-8

你的邮件无法投递。
--BOUND
Content-Type: message/delivery-status

Reporting-MTA: dns; qq.com
Final-Recipient: rfc822; HR@OldCompany.com
Action: failed
Status: 5.1.1
Diagnostic-Code: smtp; 550 5.1.1 User unknown

--BOUND
Content-Type: text/rfc822-headers

Subject: test
--BOUND--
""".encode("utf-8")


def _parse(raw):
    from email import policy
    from email.parser import BytesParser
    return BytesParser(policy=policy.default).parsebytes(raw)


def test_parse_dsn():
    from engine.bounce import parse_dsn, looks_like_bounce
    msg = _parse(DSN_RAW)
    assert looks_like_bounce(str(msg["From"]), str(msg["Subject"]))
    rows = parse_dsn(msg)
    assert len(rows) == 1, rows
    assert rows[0]["email"] == "hr@oldcompany.com"
    assert rows[0]["status"] == "5.1.1"
    assert "550" in rows[0]["reason"]


def test_parse_bounce_text_fallback():
    from engine.bounce import parse_bounce_message
    raw = """From: postmaster@163.com
Subject: 无法投递
Content-Type: text/plain; charset=utf-8

您的邮件无法投递到 dead@nowhere.com，请检查地址。
""".encode("utf-8")
    rows = parse_bounce_message(_parse(raw), "me@163.com")
    assert [r["email"] for r in rows] == ["dead@nowhere.com"], rows


def test_scan_inbox_with_fake_conn():
    from engine.bounce import scan_inbox
    from config import EmailConfig

    headers = {
        b"1": b"From: HR Wang <hr@good.com>\nSubject: Interview invitation\nDate: Wed, 21 Oct 2026 09:00:00 +0800\n\n",
        b"2": DSN_RAW[:600],
        b"3": b"From: Newsletter <news@spam.com>\nSubject: weekly\n\n",
    }

    class FakeConn:
        def select(self, *a, **k):
            return ("OK", [b"1"])

        def search(self, charset, crit):
            return ("OK", [b"1 2 3"])

        def fetch(self, uid, what):
            body = DSN_RAW if uid == b"2" else headers.get(uid, b"")
            return ("OK", [(("x (BODY %d" % len(body)).encode(), body), b")"])

        def logout(self):
            pass

    out = scan_inbox(EmailConfig(email="me@qq.com", auth_code="x"),
                     {"hr@good.com"}, days=14, conn=FakeConn())
    assert out["scanned"] == 3, out
    assert [r["email"] for r in out["replies"]] == ["hr@good.com"], out["replies"]
    assert [b["email"] for b in out["bounces"]] == ["hr@oldcompany.com"], out["bounces"]


def test_scan_detects_reply_by_message_id():
    """HR 换了邮箱回信，靠 In-Reply-To 也能认出来"""
    from engine.bounce import scan_inbox
    from config import EmailConfig
    hdr = (b"From: HR Assistant <hr-other@corp.com>\n"
           b"Subject: Re: \xe5\xae\x9e\xe4\xb9\xa0\xe7\x94\xb3\xe8\xaf\xb7\n"
           b"In-Reply-To: <abc123@qq.com>\n"
           b"Date: Wed, 21 Oct 2026 11:00:00 +0800\n\n")

    class FakeConn:
        def select(self, *a, **k): return ("OK", [b"1"])
        def search(self, charset, crit): return ("OK", [b"9"])
        def fetch(self, uid, what):
            return ("OK", [(("x (BODY %d" % len(hdr)).encode(), hdr), b")"])
        def logout(self): pass

    out = scan_inbox(EmailConfig(email="me@qq.com", auth_code="x"),
                     {"hr@target.com"}, conn=FakeConn(),
                     msgids={"<abc123@qq.com>": "hr@target.com"})
    assert [r["email"] for r in out["replies"]] == ["hr@target.com"], out["replies"]


def test_custom_provider_needs_base_url():
    import engine.ai_filler as af
    try:
        af.ai_chat("自定义（OpenAI 兼容）", "k", "m", [{"role": "user", "content": "hi"}])
    except ValueError as e:
        assert "Base URL" in str(e), str(e)
    else:
        raise AssertionError("自定义提供商没填 Base URL 必须报错")


def test_ai_test_returns_latency():
    import engine.ai_filler as af

    class Resp:
        status_code = 200
        def json(self):
            return {"choices": [{"message": {"content": "ok"}}]}

    captured = {}
    orig = af.httpx.post
    af.httpx.post = lambda url, headers=None, json=None, timeout=None: (
        captured.update({"url": url}) or Resp())
    try:
        out = af.ai_test("DeepSeek", "k", "deepseek-chat")
        assert out["reply"] == "ok" and out["latency_ms"] >= 0
        assert captured["url"].endswith("/chat/completions"), captured["url"]
    finally:
        af.httpx.post = orig


def test_vision_payload_openai_and_claude():
    """截图提取：两家 API 的消息结构必须各自正确（不联网）"""
    import json as _json
    import engine.ai_filler as af

    captured = {}

    class Resp:
        status_code = 200
        def __init__(self, payload): self.payload = payload
        def json(self): return self.payload

    def fake_openai(url, headers=None, json=None, timeout=None):
        captured["url"], captured["json"] = url, json
        return Resp({"choices": [{"message": {"content":
            '[{"company":"A","email":"a@a.com","vars":{"岗位":"量化"}}]'}}]})

    def fake_claude(url, headers=None, json=None, timeout=None):
        captured["url"], captured["json"] = url, json
        return Resp({"content": [{"text": '[{"company":"B","email":"b@b.com"}]'}]})

    orig = af.httpx.post
    try:
        af.httpx.post = fake_openai
        rows = af.ai_extract_images("ChatGPT", "k", "gpt-4o", [(b"img", "image/png")], "只看邮箱")
        assert rows[0]["company"] == "A" and rows[0]["vars"]["岗位"] == "量化"
        parts = captured["json"]["messages"][0]["content"]
        assert parts[1]["type"] == "image_url"
        assert parts[1]["image_url"]["url"].startswith("data:image/png;base64,")
        assert "只看邮箱" in captured["json"]["messages"][0]["content"][0]["text"]

        af.httpx.post = fake_claude
        rows = af.ai_extract_images("Claude", "k", "claude-sonnet-4", [(b"img", "image/jpeg")])
        assert rows[0]["email"] == "b@b.com"
        blocks = captured["json"]["messages"][0]["content"]
        assert blocks[1]["type"] == "image"
        assert blocks[1]["source"]["media_type"] == "image/jpeg"
        assert captured["json"]["max_tokens"] == 4096
    finally:
        af.httpx.post = orig


def test_normalize_rows_accepts_dict_and_list():
    from engine.ai_filler import _normalize_rows
    assert _normalize_rows({"company": "A", "email": "a@a.com"})[0]["company"] == "A"
    rows = _normalize_rows([{"company": "B", "position": "行研"}])
    assert rows[0]["vars"]["position"] == "行研"


def test_delivery_store_priority():
    import tempfile
    import engine.delivery as d

    old = d.FILE
    d.FILE = Path(tempfile.mkdtemp()) / "delivery.json"
    try:
        d.apply_scan([{"email": "dead@x.com", "reason": "550", "status": "5.1.1",
                       "date": "2026-10-20"}],
                     [{"email": "hr@good.com", "subject": "面试邀请", "date": "2026-10-21"}])
        assert "dead@x.com" in d.dead_emails()
        logs = [{"success": True, "email": "hr@good.com", "company": "好公司",
                 "time": "2026-10-19T10:00:00", "mode": "send"}]
        view = {r["email"]: r for r in d.contacts_view(logs)}
        assert view["hr@good.com"]["status"] == "replied"
        assert view["dead@x.com"]["status"] == "bounced"
        # 退信后又收到回复，不应再算作退信
        d.apply_scan([], [{"email": "dead@x.com", "subject": "误报", "date": "2026-10-22"}])
        assert "dead@x.com" not in d.dead_emails()
    finally:
        d.FILE = old


def test_check_address():
    from engine.sender import check_address
    assert check_address("hr@company.com") == ""
    assert check_address("hr@company.com.cn") == ""
    assert check_address("") == "这一行没有邮箱"
    assert "格式" in check_address("not-an-email")
    assert "多个" in check_address("a@x.com,b@y.com")
    assert "多个" in check_address("a@x.com；b@y.com")
    assert "格式" in check_address("a@x")          # 没有顶级域
    assert "只填邮箱" in check_address("张三 <a@x.com>")  # 带姓名没拆开


def test_send_batch_rejects_bad_addresses():
    """坏地址不能占额度、不能进 SMTP 循环"""
    class FakeLink:
        sent = []
        def __init__(self, cfg): pass
        def send(self, msg): FakeLink.sent.append(msg["To"])
        def close(self): pass

    from config import EmailConfig
    orig = sender._SmtpLink
    sender._SmtpLink = FakeLink
    try:
        recs = [{"company": "好", "email": "ok@x.com"},
                {"company": "空", "email": ""},
                {"company": "坏", "email": "not-an-email"},
                {"company": "多", "email": "a@x.com,b@y.com"}]
        out = sender.send_batch(EmailConfig(email="me@qq.com"), recs, "hi", "s", [],
                                send_mode="send", rate=600)
        got = [(r["status"], r["company"]) for r in out["results"]]
        assert got == [("sent", "好"), ("failed", "空"), ("failed", "坏"), ("failed", "多")], got
        assert FakeLink.sent == ["ok@x.com"], FakeLink.sent
    finally:
        sender._SmtpLink = orig


def test_send_batch_skips_flagged_rows():
    class FakeLink:
        sent = []

        def __init__(self, cfg):
            pass

        def send(self, msg):
            FakeLink.sent.append(msg["To"])

        def close(self):
            pass

    from config import EmailConfig
    orig = sender._SmtpLink
    sender._SmtpLink = FakeLink
    try:
        recs = [{"company": "A", "email": "a@x.com"},
                {"company": "B", "email": "b@x.com", "skip_reason": "退信地址，已自动跳过"}]
        out = sender.send_batch(EmailConfig(email="me@qq.com"), recs, "hi", "s", [],
                                send_mode="send", rate=600)
        assert [r["status"] for r in out["results"]] == ["sent", "skipped"], out["results"]
        assert FakeLink.sent == ["a@x.com"]
        assert "退信" in out["results"][1]["error"]
    finally:
        sender._SmtpLink = orig


# ── 文件读取 / SOUL / 新端点 ──

def _docx_bytes(text):
    """手搓一个最小 docx（zip + word/document.xml），避免测试依赖 python-docx"""
    import zipfile
    buf = io.BytesIO()
    paras = "".join(f"<w:p><w:r><w:t>{line}</w:t></w:r></w:p>" for line in text.split("\n"))
    xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           f"<w:body>{paras}</w:body></w:document>")
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", xml)
    return buf.getvalue()


def test_file_reader_docx_xlsx_txt():
    import tempfile
    from engine.file_reader import read_text, file_kind

    tmp = Path(tempfile.mkdtemp())
    docx = tmp / "招聘.docx"
    docx.write_bytes(_docx_bytes("华泰证券招聘\n简历发 hr@htsc.com"))
    assert file_kind(docx) == "docx"
    got = read_text(docx)
    assert "华泰证券" in got and "hr@htsc.com" in got, got

    xlsx = tmp / "名单.xlsx"
    xlsx.write_bytes(_xlsx([["公司", "邮箱"], ["中信", "hr@citics.com"]]))
    sheet = read_text(xlsx)
    assert "中信" in sheet and "hr@citics.com" in sheet, sheet

    gbk = tmp / "旧编码.txt"
    gbk.write_bytes("公司：国泰君安 hr@gtja.com".encode("gb18030"))
    assert "国泰君安" in read_text(gbk), "GBK 文本应能读出来"

    img = tmp / "截图.png"
    img.write_bytes(b"\x89PNG\r\n\x1a\n")
    assert file_kind(img) == "image" and read_text(img) == ""


def test_soul_prompts_exist():
    from engine.ai_filler import load_soul
    recipient = load_soul("recipient")
    assert "JSON" in recipient and "宁可漏" in recipient, "收件人 soul 应含关键约束"
    assert "company" in recipient and "email" in recipient
    resume = load_soul("resume")
    assert "姓名" in resume and "学校" in resume


def test_parse_files_endpoint_plumbing():
    from fastapi.testclient import TestClient
    import server

    orig = server.ai_extract_files
    server.ai_extract_files = lambda *a, **k: [
        {"company": "华泰证券", "email": "hr@htsc.com", "subject": "",
         "vars": {"岗位": "量化"}, "note": ""}
    ]
    try:
        client = TestClient(server.app)
        r = client.post("/api/ai/parse-files",
                        files=[("files", ("招聘.txt", "华泰证券 hr@htsc.com".encode(), "text/plain"))],
                        data={"provider": "DeepSeek", "model": "deepseek-flash", "api_key": "k"})
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["recipients"][0]["email"] == "hr@htsc.com"
        assert d["files_used"] and "招聘.txt" in d["files_used"][0]

        # 空文件也要给出人话错误
        empty = client.post("/api/ai/parse-files",
                            files=[("files", ("空白.docx", _docx_bytes(""), "application/octet-stream"))],
                            data={"provider": "DeepSeek", "model": "m", "api_key": "k"})
        assert empty.status_code == 400, empty.text
    finally:
        server.ai_extract_files = orig


def test_resume_fill_endpoint():
    import tempfile
    from fastapi.testclient import TestClient
    import server

    tmp = Path(tempfile.mkdtemp())
    (tmp / "简历.txt").write_text("张三 华东理工大学 金融学 27届", encoding="utf-8")
    orig_dir, orig_ai = server.ATTACH_DIR, server.ai_parse_resume
    server.ATTACH_DIR = tmp
    server.ai_parse_resume = lambda *a, **k: {
        "姓名": "张三", "学校": "华东理工大学", "专业": "金融学",
        "年级": "27届", "电话": "", "邮箱": "", "note": "",
    }
    try:
        client = TestClient(server.app)
        ok = client.post("/api/ai/resume-fill", json={
            "filename": str(tmp / "简历.txt"), "provider": "DeepSeek",
            "model": "m", "api_key": "k"})
        assert ok.status_code == 200, ok.text
        assert ok.json()["fields"]["学校"] == "华东理工大学"

        # 目录外的路径必须拒绝
        bad = client.post("/api/ai/resume-fill", json={
            "filename": "/etc/passwd", "provider": "DeepSeek", "model": "m", "api_key": "k"})
        assert bad.status_code == 400, bad.text
    finally:
        server.ATTACH_DIR, server.ai_parse_resume = orig_dir, orig_ai


def test_slice_batch_first_and_rest():
    """先发第一封 / 发剩下：limit 截断，skip_sent 按成功记录跳过"""
    import server
    rows = [{"company": "A", "email": "a@x.com"}, {"company": "B", "email": "b@x.com"},
            {"company": "C", "email": "c@x.com"}]
    logs = [{"email": "A@X.com", "success": True, "status": "sent"},
            {"email": "b@x.com", "success": False, "status": "failed"},
            {"email": "c@x.com", "success": True, "status": "draft"}]

    # 一封都没发：limit=1 只给第一封
    got, skipped = server.slice_batch(rows, skip_sent=True, limit=1, logs=[])
    assert [r["email"] for r in got] == ["a@x.com"], got
    assert skipped == 0, skipped

    # 已发过 a：limit=1 应该是“下一个没发过的”b，而不是重发 a
    got, skipped = server.slice_batch(rows, skip_sent=True, limit=1, logs=logs)
    assert [r["email"] for r in got] == ["b@x.com"], got
    assert skipped == 1, skipped

    # 发过 1 封（大小写不同也算）：剩下 2 封，草稿/失败的不算已发
    got, skipped = server.slice_batch(rows, skip_sent=True, limit=0, logs=logs)
    assert [r["email"] for r in got] == ["b@x.com", "c@x.com"], got
    assert skipped == 1, skipped

    # 全发过：空列表，调用方据此报错
    got, skipped = server.slice_batch(rows, skip_sent=True, limit=0,
                                      logs=[{"email": r["email"], "success": True, "status": "sent"} for r in rows])
    assert got == [] and skipped == 3, (got, skipped)

    # 不跳过时 limit=1 就是原样第一封
    got, skipped = server.slice_batch(rows, skip_sent=False, limit=1, logs=logs)
    assert len(got) == 1 and skipped == 0


def test_recent_logs_window():
    """投递记录只取最近 30 天；时间解析不了的保留，31 天前的丢掉"""
    import server
    from datetime import datetime, timedelta
    now = datetime.now()
    logs = [
        {"email": "now@x.com", "time": now.isoformat(timespec="seconds"), "success": True, "status": "sent"},
        {"email": "edge@x.com", "time": (now - timedelta(days=29, hours=23)).isoformat(timespec="seconds"),
         "success": True, "status": "sent"},
        {"email": "old@x.com", "time": (now - timedelta(days=31)).isoformat(timespec="seconds"),
         "success": True, "status": "sent"},
        {"email": "bad@x.com", "time": "not-a-date", "success": True, "status": "sent"},
    ]
    got = [e["email"] for e in server.recent_logs(logs=logs)]
    assert "now@x.com" in got and "edge@x.com" in got, got
    assert "old@x.com" not in got, got
    assert "bad@x.com" in got, got          # 解析不了就保留，不静默丢数据


def test_safe_filename_cross_platform():
    """附件名要能落到 Windows 上：非法字符、设备名、反斜杠路径、超长名"""
    import server
    assert server.safe_filename("简历:2026*.pdf") == "简历_2026_.pdf"
    assert server.safe_filename("CON.pdf") == "_CON.pdf"
    assert server.safe_filename(r"C:\Users\me\cv.docx") == "cv.docx"
    assert server.safe_filename("   ") == "attachment"
    assert server.safe_filename("") == "attachment"
    long_name = server.safe_filename("长" * 200 + ".pdf")
    assert len(long_name) <= 90 and long_name.endswith(".pdf")
    # 正常中文名必须原样保留
    assert server.safe_filename("戚俊皓-华东理工大学-本科-金融.pdf") == "戚俊皓-华东理工大学-本科-金融.pdf"


def test_default_body_template_contents():
    """默认正文就是用户给的那封：段落、变量、落款都要在"""
    import config
    t = config.DEFAULT_BODY_TEMPLATE
    assert t.startswith("尊敬的招聘负责人：")
    assert "我是{姓名}，{学校}{专业}专业{年级}学生（2027年毕业），现申请贵公司的实习岗位。" in t
    assert "我具备扎实的金融理论基础，熟练使用Wind、Choice等金融数据终端及Excel数据处理工具" in t
    assert "随信附上我的简历，敬请查阅。期待您的回复！" in t
    assert "此致\n敬礼" in t
    assert t.rstrip().endswith("{姓名}\n{邮箱}")
    assert "XXX" not in t            # 占位符已经换成变量，不会发出「我是XXX」
    assert "「" not in t and "」" not in t


if __name__ == "__main__":
    tests = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_")]
    failed = 0
    for name, fn in tests:
        try:
            fn()
            print(f"  ok  {name}")
        except Exception as e:
            failed += 1
            print(f"FAIL  {name}: {type(e).__name__}: {e}")
    print(f"\n{len(tests) - failed}/{len(tests)} passed")
    sys.exit(1 if failed else 0)
